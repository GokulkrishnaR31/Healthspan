import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation in 16:9 Widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette Constants
PURPLE_PRIMARY = RGBColor(88, 37, 131)     # #582583 Official REC Purple
PURPLE_LIGHT = RGBColor(246, 243, 250)      # #F6F3FA Card Background
PURPLE_BORDER = RGBColor(215, 203, 230)     # #D7CBE6 Card Border
PURPLE_ACCENT = RGBColor(123, 63, 175)     # #7B3FAF Badge Purple
TEXT_DARK = RGBColor(35, 35, 45)            # #23232D Primary Text
TEXT_MUTED = RGBColor(100, 100, 115)        # #646473 Secondary Text
TEXT_WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(248, 247, 252)
TABLE_HEADER = RGBColor(88, 37, 131)
TABLE_ALT = RGBColor(248, 246, 252)

FONT_FAMILY = "Segoe UI"
BLANK_LAYOUT = prs.slide_layouts[6]
IMG_DIR = r"g:\rproject\output"

def add_header(slide, title_text, slide_num):
    # Top Purple Header Banner
    banner = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(13.333), Inches(1.1)
    )
    banner.fill.solid()
    banner.fill.fore_color.rgb = PURPLE_PRIMARY
    banner.line.color.rgb = PURPLE_PRIMARY

    # Title Text inside Banner
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(12.333), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.name = FONT_FAMILY
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Slide Number at bottom right
    if slide_num > 1:
        numBox = slide.shapes.add_textbox(Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.35))
        np = numBox.text_frame.paragraphs[0]
        np.text = str(slide_num)
        np.font.name = FONT_FAMILY
        np.font.size = Pt(11)
        np.font.color.rgb = TEXT_MUTED
        np.alignment = PP_ALIGN.RIGHT

def create_card(slide, left, top, width, height, bg_color=PURPLE_LIGHT, border_color=PURPLE_BORDER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

# ==========================================
# SLIDE 1: TITLE SLIDE
# ==========================================
s1 = prs.slides.add_slide(BLANK_LAYOUT)

# College Header Text
c_box = s1.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(1.4))
tf1 = c_box.text_frame
p1 = tf1.paragraphs[0]
p1.text = "RAJALAKSHMI ENGINEERING COLLEGE"
p1.font.name = FONT_FAMILY
p1.font.size = Pt(26)
p1.font.bold = True
p1.font.color.rgb = PURPLE_PRIMARY
p1.alignment = PP_ALIGN.CENTER

p2 = tf1.add_paragraph()
p2.text = "Approved by AICTE | Affiliated to Anna University | Accredited by NAAC\nDepartment of Information Technology"
p2.font.name = FONT_FAMILY
p2.font.size = Pt(13)
p2.font.color.rgb = TEXT_MUTED
p2.alignment = PP_ALIGN.CENTER

# Course Badge
p3 = tf1.add_paragraph()
p3.text = "IT23721 – DATA SCIENCE USING R"
p3.font.name = FONT_FAMILY
p3.font.size = Pt(14)
p3.font.bold = True
p3.font.color.rgb = PURPLE_ACCENT
p3.alignment = PP_ALIGN.CENTER

# Main Project Title Box
t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.8), Inches(11.333), Inches(2.2))
tf_t = t_box.text_frame
tp1 = tf_t.paragraphs[0]
tp1.text = "AIRSIGHT — Large-Scale Air Quality Index (AQI)\nAnalytics & Predictive Forecasting Platform"
tp1.font.name = FONT_FAMILY
tp1.font.size = Pt(28)
tp1.font.bold = True
tp1.font.color.rgb = TEXT_DARK
tp1.alignment = PP_ALIGN.CENTER

tp2 = tf_t.add_paragraph()
tp2.text = "Spatio-Temporal Analysis, Inferential Statistics & Machine Learning in R"
tp2.font.name = FONT_FAMILY
tp2.font.size = Pt(15)
tp2.font.italic = True
tp2.font.color.rgb = PURPLE_PRIMARY
tp2.alignment = PP_ALIGN.CENTER

# Author Details Box
auth_box = s1.shapes.add_textbox(Inches(7.5), Inches(5.6), Inches(5.0), Inches(1.2))
tf_a = auth_box.text_frame
ap1 = tf_a.paragraphs[0]
ap1.text = "2116231001054\n[Your Name]\nB.Tech Information Technology"
ap1.font.name = FONT_FAMILY
ap1.font.size = Pt(14)
ap1.font.bold = True
ap1.font.color.rgb = TEXT_DARK
ap1.alignment = PP_ALIGN.RIGHT

# ==========================================
# SLIDE 2: ABSTRACT
# ==========================================
s2 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s2, "Abstract", 2)

# Abstract Bullets
ab_box = s2.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(3.4))
tf_ab = ab_box.text_frame
tf_ab.word_wrap = True

bullets_s2 = [
    "Air quality deterioration poses severe environmental and public health hazards across Indian metropolitan and industrial zones.",
    "AIRSIGHT is an enterprise R-based analytics platform that processes, visualizes, statistically tests, and predicts multi-year AQI trends.",
    "Dataset: 235,785 national environmental observations across 32 States and Union Territories from 2022 to 2025.",
    "Engineered 16 ggplot2 analytical charts, conducted Welch's t-test and 3 One-Way ANOVAs, and trained multiple ML regression models.",
    "Random Forest achieved the best prediction performance: RMSE 48.28, MAE 34.46, and R² 0.5412.",
    "Deployed as an interactive 8-tab R Shiny dashboard with Plumber REST API endpoints for real-time monitoring."
]
for b in bullets_s2:
    p = tf_ab.add_paragraph() if tf_ab.paragraphs[0].text else tf_ab.paragraphs[0]
    p.text = "•  " + b
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(10)

# Bottom Stat Cards
kpis_s2 = [
    ("235,785", "raw observation records"),
    ("16 Charts", "ggplot2 visualizations"),
    ("32 States", "spatio-temporal coverage"),
    ("R Shiny & API", "interactive dashboard")
]
card_w = 2.7
for i, (val, label) in enumerate(kpis_s2):
    left = 0.8 + i * 2.95
    create_card(s2, left, 5.2, card_w, 1.4)
    cbox = s2.shapes.add_textbox(Inches(left), Inches(5.3), Inches(card_w), Inches(1.2))
    ctf = cbox.text_frame
    cp1 = ctf.paragraphs[0]
    cp1.text = val
    cp1.font.name = FONT_FAMILY
    cp1.font.size = Pt(22)
    cp1.font.bold = True
    cp1.font.color.rgb = PURPLE_PRIMARY
    cp1.alignment = PP_ALIGN.CENTER
    
    cp2 = ctf.add_paragraph()
    cp2.text = label
    cp2.font.name = FONT_FAMILY
    cp2.font.size = Pt(11)
    cp2.font.color.rgb = TEXT_MUTED
    cp2.alignment = PP_ALIGN.CENTER

# ==========================================
# SLIDE 3: INTRODUCTION
# ==========================================
s3 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s3, "Introduction", 3)

in_box = s3.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(3.4))
tf_in = in_box.text_frame
tf_in.word_wrap = True

bullets_s3 = [
    "Air Quality Index (AQI) standardizes complex pollutant concentrations into an easily understandable health scale (0–500).",
    "Particulate matter (PM2.5, PM10) along with gases (NO₂, SO₂, CO, O₃) are the primary determinants of respiratory vulnerability.",
    "Exploratory Data Analysis (EDA) reveals severe non-linear seasonal cycles and spatial clusters across regions.",
    "Manual analysis across massive multi-sensor CSV datasets is slow, error-prone, and lacks standardized statistical rigor.",
    "R is selected for its best-in-class statistical foundations, ggplot2 graphical pipeline, and reactive Shiny dashboarding.",
    "AIRSIGHT delivers a unified platform integrating data cleaning, hypothesis testing, machine learning, and policy simulation."
]
for b in bullets_s3:
    p = tf_in.add_paragraph() if tf_in.paragraphs[0].text else tf_in.paragraphs[0]
    p.text = "•  " + b
    p.font.name = FONT_FAMILY
    p.font.size = Pt(14)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(10)

# Process Flow Diagram at bottom
flow_steps = ["Raw Sensor CSV", "Tidy Wrangling", "Inferential EDA", "ML Prediction", "R Shiny Dashboard"]
for i, step in enumerate(flow_steps):
    left = 0.8 + i * 2.38
    create_card(s3, left, 5.3, 2.1, 1.2, bg_color=PURPLE_LIGHT, border_color=PURPLE_PRIMARY)
    box = s3.shapes.add_textbox(Inches(left), Inches(5.5), Inches(2.1), Inches(0.8))
    p = box.text_frame.paragraphs[0]
    p.text = f"{i+1}. {step}"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PURPLE_PRIMARY
    p.alignment = PP_ALIGN.CENTER
    if i < 4:
        arr_box = s3.shapes.add_textbox(Inches(left + 2.05), Inches(5.6), Inches(0.35), Inches(0.6))
        ap = arr_box.text_frame.paragraphs[0]
        ap.text = "➔"
        ap.font.size = Pt(16)
        ap.font.color.rgb = PURPLE_ACCENT

# ==========================================
# SLIDE 4: PROBLEM STATEMENT
# ==========================================
s4 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s4, "Problem Statement", 4)

ps_box = s4.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(7.8), Inches(5.2))
tf_ps = ps_box.text_frame
tf_ps.word_wrap = True

ps_items = [
    ("Sensor Noise & Incompleteness", "Environmental monitoring logs suffer from missing readings, sensor calibration shifts, and multi-gas discrepancies that corrupt baseline analysis."),
    ("Non-Linear Weather Dynamics", "Pollutant dispersion is heavily driven by seasonal atmospheric inversions, temperature dips in winter, and precipitation washout in monsoons."),
    ("Massive Regional Inequality", "Vast disparities exist between severe industrial zones (Indo-Gangetic plains) and pristine coastal regions, requiring localized modeling."),
    ("Lack of Automated Simulation", "City planners lack interactive tools to forecast severe pollution spikes and simulate the real-time impact of emission reduction policies.")
]
for title, desc in ps_items:
    p1 = tf_ps.add_paragraph() if tf_ps.paragraphs[0].text else tf_ps.paragraphs[0]
    p1.text = title + ":"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = PURPLE_PRIMARY
    
    p2 = tf_ps.add_paragraph()
    p2.text = desc
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_DARK
    p2.space_after = Pt(12)

# Right Badges
r_kpis = [
    ("44.5%", "Non-Compliant Days\n(CPCB >100 AQI Standard)"),
    ("82.2%", "Exceeding WHO Limits\n(Daily AQI > 50 Safe Level)"),
    ("1,125", "Extreme Pollution Anomaly\nEvents Detected")
]
for i, (val, label) in enumerate(r_kpis):
    top = 1.5 + i * 1.75
    create_card(s4, 9.0, top, 3.5, 1.45)
    rbox = s4.shapes.add_textbox(Inches(9.0), Inches(top + 0.1), Inches(3.5), Inches(1.25))
    rtf = rbox.text_frame
    rp1 = rtf.paragraphs[0]
    rp1.text = val
    rp1.font.name = FONT_FAMILY
    rp1.font.size = Pt(24)
    rp1.font.bold = True
    rp1.font.color.rgb = PURPLE_PRIMARY
    rp1.alignment = PP_ALIGN.CENTER
    
    rp2 = rtf.add_paragraph()
    rp2.text = label
    rp2.font.name = FONT_FAMILY
    rp2.font.size = Pt(11)
    rp2.font.color.rgb = TEXT_MUTED
    rp2.alignment = PP_ALIGN.CENTER

# ==========================================
# SLIDE 5: OBJECTIVES
# ==========================================
s5 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s5, "Objectives", 5)

# Left Column
obj_l = [
    "1. Load 235,785 environmental records with schema validation.",
    "2. Clean missing data, dates, and standardize pollutant categories.",
    "3. Engineer features: month_num, is_winter, is_monsoon, pollutant_score.",
    "4. Build 16 publication-grade ggplot2 exploratory visualizations.",
    "5. Perform Welch's t-test and 3 One-Way ANOVAs (State, Month, Pollutant)."
]
# Right Column
obj_r = [
    "6. Calculate CPCB and WHO national & global compliance benchmarks.",
    "7. Train and compare Linear Regression vs Random Forest (80/20 split).",
    "8. Implement a 7-day ARIMA / time-series forecasting model.",
    "9. Develop a Health Policy & Air Purifier CADR impact simulator.",
    "10. Deploy an 8-tab interactive R Shiny dashboard and Plumber REST API."
]

ol_box = s5.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
tf_ol = ol_box.text_frame
tf_ol.word_wrap = True
for item in obj_l:
    p = tf_ol.add_paragraph() if tf_ol.paragraphs[0].text else tf_ol.paragraphs[0]
    p.text = item
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(16)

or_box = s5.shapes.add_textbox(Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.2))
tf_or = or_box.text_frame
tf_or.word_wrap = True
for item in obj_r:
    p = tf_or.add_paragraph() if tf_or.paragraphs[0].text else tf_or.paragraphs[0]
    p.text = item
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(16)

# ==========================================
# SLIDE 6: SCOPE
# ==========================================
s6 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s6, "Scope", 6)

scope_cards = [
    ("Architecture & Pipeline", [
        "Primary statistical core in R 4.4 with 13 modular scripts.",
        "Interactive R Shiny dashboard (8 tabs) + Plumber REST API.",
        "React + Vite frontend for cross-platform public portal."
    ]),
    ("Statistical & ML Analysis", [
        "Descriptive metrics (Mean, Median, SD, IQR, Outliers).",
        "3x One-Way ANOVA tests across states, months, and pollutants.",
        "Random Forest regression predicting AQI (RMSE 48.28, R² 0.54)."
    ]),
    ("Practical Health Utility", [
        "AQLI life expectancy lost calculation (years/months lost).",
        "Air purifier CADR sizing and indoor filtration estimator.",
        "Municipal policy simulator estimating avoided ER visits."
    ]),
    ("Dataset Scope & Boundaries", [
        "Covers 32 Indian States/UTs across multi-year timeline.",
        "Analyzes standard criteria pollutants: PM2.5, PM10, NO₂, SO₂, CO, O₃.",
        "Static benchmark validated against live real-time API feeds."
    ])
]

for idx, (title, bullets) in enumerate(scope_cards):
    r = idx // 2
    c = idx % 2
    left = 0.8 + c * 5.95
    top = 1.5 + r * 2.7
    create_card(s6, left, top, 5.7, 2.5)
    
    tbox = s6.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.15), Inches(5.3), Inches(2.2))
    ttf = tbox.text_frame
    ttf.word_wrap = True
    tp = ttf.paragraphs[0]
    tp.text = title
    tp.font.name = FONT_FAMILY
    tp.font.size = Pt(14)
    tp.font.bold = True
    tp.font.color.rgb = PURPLE_PRIMARY
    tp.space_after = Pt(6)
    
    for b in bullets:
        bp = ttf.add_paragraph()
        bp.text = "•  " + b
        bp.font.name = FONT_FAMILY
        bp.font.size = Pt(12)
        bp.font.color.rgb = TEXT_DARK
        bp.space_after = Pt(4)

# ==========================================
# SLIDE 7: LITERATURE REVIEW
# ==========================================
s7 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s7, "Literature Review", 7)

# Literature Table
rows, cols = 4, 4
left, top, width, height = Inches(0.8), Inches(1.4), Inches(11.733), Inches(3.0)
table_shape = s7.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table
table.columns[0].width = Inches(2.2)
table.columns[1].width = Inches(2.7)
table.columns[2].width = Inches(3.2)
table.columns[3].width = Inches(3.633)

headers = ["Author & Year", "Focus", "Method Used", "Limitation / Proposed Advantage"]
for c, h in enumerate(headers):
    cell = table.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = TABLE_HEADER
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

lit_data = [
    ("Goyal et al. (2020)", "Urban AQI Assessment", "Univariate averages & trendlines", "Static analysis; no predictive machine learning."),
    ("Kumar et al. (2022)", "Pollutant Inter-correlation", "Linear OLS regression", "Lacks seasonal controls and interactive UI."),
    ("AIRSIGHT (2026)", "Full-Stack AQI Platform", "R, Random Forest, Shiny, Plumber", "Overcomes gaps with 8-tab dashboard, ANOVA & ML.")
]
for r, row in enumerate(lit_data):
    for c, val in enumerate(row):
        cell = table.cell(r + 1, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_ALT if r % 2 == 1 else RGBColor(255, 255, 255)
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = (r == 2)
        p.font.color.rgb = PURPLE_PRIMARY if r == 2 else TEXT_DARK

# Bottom Takeaways
t_box = s7.shapes.add_textbox(Inches(0.8), Inches(4.7), Inches(11.733), Inches(2.2))
ttf = t_box.text_frame
ttf.word_wrap = True
tp1 = ttf.paragraphs[0]
tp1.text = "Existing Approaches vs AIRSIGHT"
tp1.font.name = FONT_FAMILY
tp1.font.size = Pt(14)
tp1.font.bold = True
tp1.font.color.rgb = PURPLE_PRIMARY
tp1.space_after = Pt(6)

bullets_lit = [
    "Traditional systems rely on ad-hoc scripts to produce isolated, non-reproducible static charts.",
    "Prior studies rarely combine inferential hypothesis testing (ANOVA) with machine learning ensembles.",
    "AIRSIGHT provides end-to-end automation: ingestion ➔ cleaning ➔ hypothesis testing ➔ ML ➔ interactive Shiny UI."
]
for b in bullets_lit:
    bp = ttf.add_paragraph()
    bp.text = "•  " + b
    bp.font.name = FONT_FAMILY
    bp.font.size = Pt(12)
    bp.font.color.rgb = TEXT_DARK
    bp.space_after = Pt(4)

# ==========================================
# SLIDE 8: SYSTEM REQUIREMENTS
# ==========================================
s8 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s8, "System Requirements", 8)

# Left: Hardware Table
h_table_shape = s8.shapes.add_table(5, 3, Inches(0.8), Inches(1.5), Inches(5.6), Inches(2.8))
ht = h_table_shape.table
ht.columns[0].width = Inches(1.6)
ht.columns[1].width = Inches(1.8)
ht.columns[2].width = Inches(2.2)

for c, h in enumerate(["Hardware", "Minimum", "Recommended"]):
    cell = ht.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = TABLE_HEADER
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

hw_rows = [
    ("Processor", "Core i3 (2.0 GHz)", "Core i5 / i7 / Ryzen 5+"),
    ("RAM", "4 GB", "8 GB to 16 GB"),
    ("Storage", "1 GB Free Space", "4 GB Free SSD"),
    ("Display", "1024 x 768", "1920 x 1080 (Full HD)")
]
for r, row in enumerate(hw_rows):
    for c, val in enumerate(row):
        cell = ht.cell(r + 1, c)
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK

# Right: R Packages Table
p_table_shape = s8.shapes.add_table(8, 2, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
pt = p_table_shape.table
pt.columns[0].width = Inches(2.2)
pt.columns[1].width = Inches(3.5)

for c, h in enumerate(["Package", "Purpose & Utility"]):
    cell = pt.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = TABLE_HEADER
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

pkg_rows = [
    ("shiny / shinydashboard", "Interactive multi-tab reactive web dashboard"),
    ("tidyverse / dplyr", "Data wrangling, cleaning & pipeline transformations"),
    ("ggplot2 / plotly", "16 publication-quality statistical visualizations"),
    ("randomForest", "Non-linear ensemble predictive regression model"),
    ("corrplot", "Correlation matrix heatmaps & collinearity analysis"),
    ("leaflet", "Geospatial interactive mapping of station AQI"),
    ("plumber", "REST API backend serving JSON endpoints")
]
for r, row in enumerate(pkg_rows):
    for c, val in enumerate(row):
        cell = pt.cell(r + 1, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_ALT if r % 2 == 1 else RGBColor(255, 255, 255)
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK

# Bottom left notes
h_box = s8.shapes.add_textbox(Inches(0.8), Inches(4.6), Inches(5.6), Inches(2.2))
htf = h_box.text_frame
htf.word_wrap = True
hp1 = htf.paragraphs[0]
hp1.text = "Environment & Tools Used"
hp1.font.name = FONT_FAMILY
hp1.font.size = Pt(13)
hp1.font.bold = True
hp1.font.color.rgb = PURPLE_PRIMARY
hp1.space_after = Pt(4)

env_bullets = [
    "OS: Windows 10/11, macOS or Linux compatible.",
    "R Version: R 4.4.2 / RStudio IDE (Tested).",
    "Frontend Engine: React + Vite with Chart.js & Leaflet.",
    "Backend Engine: Plumber REST API on port 8000."
]
for b in env_bullets:
    p = htf.add_paragraph()
    p.text = "•  " + b
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(3)

# ==========================================
# SLIDE 9: DATASET DESCRIPTION
# ==========================================
s9 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s9, "Dataset Description", 9)

# Left Stat Badges (6 grid)
d_stats = [
    ("235,785", "total records"),
    ("32", "states & UTs"),
    ("4 Years", "2022 – 2025"),
    ("CPCB", "official source"),
    ("6 Criteria", "PM2.5, PM10, gases"),
    ("Clean RDS", "data objects")
]
for idx, (val, label) in enumerate(d_stats):
    r = idx // 2
    c = idx % 2
    left = 0.8 + c * 2.2
    top = 1.5 + r * 1.7
    create_card(s9, left, top, 2.0, 1.45)
    box = s9.shapes.add_textbox(Inches(left), Inches(top + 0.1), Inches(2.0), Inches(1.25))
    tf = box.text_frame
    p1 = tf.paragraphs[0]
    p1.text = val
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = PURPLE_PRIMARY
    p1.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MUTED
    p2.alignment = PP_ALIGN.CENTER

# Right Attribute Table
d_table_shape = s9.shapes.add_table(8, 3, Inches(5.4), Inches(1.5), Inches(7.1), Inches(5.2))
dt = d_table_shape.table
dt.columns[0].width = Inches(2.0)
dt.columns[1].width = Inches(1.1)
dt.columns[2].width = Inches(4.0)

for c, h in enumerate(["Attribute", "Type", "Meaning & Domain"]):
    cell = dt.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = TABLE_HEADER
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

attr_rows = [
    ("state", "Factor", "Indian State or Union Territory (32 distinct)"),
    ("city", "Factor", "City / Urban agglomeration location"),
    ("aqi_value", "Numeric", "Air Quality Index continuous value (0 to 500)"),
    ("aqi_category", "Factor", "Good, Satisfactory, Moderate, Poor, Very Poor, Severe"),
    ("prominent_pollutants", "Factor", "Dominant pollutant (PM2.5, PM10, NO₂, SO₂, O₃, CO)"),
    ("date / month_num", "Date/Int", "Timestamp and month (1–12) for seasonal grouping"),
    ("number_of_stations", "Numeric", "Active monitoring stations reporting in city")
]
for r, row in enumerate(attr_rows):
    for c, val in enumerate(row):
        cell = dt.cell(r + 1, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_ALT if r % 2 == 1 else RGBColor(255, 255, 255)
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 10: METHODOLOGY
# ==========================================
s10 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s10, "Methodology", 10)

meth_cards = [
    ("1. Ingestion & Validation", "Load raw CSV with defensive checks. Validate headers, column types, UTF-8 encoding, and row integrity."),
    ("2. Cleaning & Engineering", "Impute missing values, parse date timestamps, and engineer features: month_num, is_winter, is_monsoon, and pollutant_score."),
    ("3. Visual Exploratory Analysis", "Construct 16 publication-grade ggplot2 visualizations covering distributions, spatial boxplots, and heatmaps."),
    ("4. Inferential Statistics", "Conduct Welch's Two-Sample t-test and 3 One-Way ANOVAs (State, Month, Pollutant) to evaluate statistical significance."),
    ("5. ML & Time-Series Modeling", "Train Multiple Linear Regression and Random Forest (200 trees) on 80/20 split; generate 7-day ARIMA forecasts."),
    ("6. Dashboard & Deployment", "Deploy full-stack interactive R Shiny dashboard and Plumber REST API with policy simulation and health risk calculators.")
]

for idx, (title, desc) in enumerate(meth_cards):
    r = idx // 3
    c = idx % 3
    left = 0.8 + c * 3.95
    top = 1.6 + r * 2.7
    create_card(s10, left, top, 3.8, 2.45)
    
    box = s10.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.15), Inches(3.5), Inches(2.15))
    tf = box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = PURPLE_PRIMARY
    p1.space_after = Pt(6)
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_DARK

# ==========================================
# SLIDE 11: SYSTEM ARCHITECTURE
# ==========================================
s11 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s11, "System Architecture", 11)

pillars = [
    ("1. Data Ingestion\n& Storage", [
        "CPCB / OpenAQ CSV",
        "235,785 observations",
        "Defensive loading",
        "Schema & type validation"
    ]),
    ("2. Data Cleaning &\nPreprocessing", [
        "NA & outlier imputation",
        "Date parsing & temporal tags",
        "Feature: is_winter, is_monsoon",
        "Feature: pollutant_score"
    ]),
    ("3. Exploratory &\nInferential Analysis", [
        "16 ggplot2 Visuals",
        "Pearson/Spearman corr",
        "3x One-Way ANOVAs",
        "Welch's Two-Sample t-test"
    ]),
    ("4. ML & Predictive\nModeling", [
        "80/20 Train-Test split",
        "Linear Regression model",
        "Random Forest (ntree=200)",
        "7-Day ARIMA forecast"
    ]),
    ("5. Output & Decision\nSupport", [
        "8-Tab R Shiny Dashboard",
        "Plumber REST API",
        "Policy Simulator (ER drop)",
        "AQLI & CADR calculator"
    ])
]

for idx, (title, items) in enumerate(pillars):
    left = 0.8 + idx * 2.38
    # Pillar Header Box
    hshape = slide = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(1.5), Inches(2.2), Inches(0.8))
    hshape.fill.solid()
    hshape.fill.fore_color.rgb = PURPLE_PRIMARY
    hshape.line.color.rgb = PURPLE_PRIMARY
    hp = hshape.text_frame.paragraphs[0]
    hp.text = title
    hp.font.name = FONT_FAMILY
    hp.font.size = Pt(11)
    hp.font.bold = True
    hp.font.color.rgb = TEXT_WHITE
    hp.alignment = PP_ALIGN.CENTER
    
    # Pillar Body Box
    bshape = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(2.4), Inches(2.2), Inches(4.0))
    bshape.fill.solid()
    bshape.fill.fore_color.rgb = PURPLE_LIGHT
    bshape.line.color.rgb = PURPLE_BORDER
    btf = bshape.text_frame
    btf.word_wrap = True
    for item in items:
        p = btf.add_paragraph() if btf.paragraphs[0].text else btf.paragraphs[0]
        p.text = "• " + item
        p.font.name = FONT_FAMILY
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

# Bottom Pipeline Flow Footer
foot_shape = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.45))
foot_shape.fill.solid()
foot_shape.fill.fore_color.rgb = PURPLE_PRIMARY
foot_shape.line.color.rgb = PURPLE_PRIMARY
fp = foot_shape.text_frame.paragraphs[0]
fp.text = "Data Ingestion ➔ Data Cleaning ➔ Feature Engineering ➔ Inferential EDA ➔ ML Training ➔ R Shiny Dashboard"
fp.font.name = FONT_FAMILY
fp.font.size = Pt(11)
fp.font.bold = True
fp.font.color.rgb = TEXT_WHITE
fp.alignment = PP_ALIGN.CENTER

# ==========================================
# HELPER: DUAL-CHART VISUALIZATION SLIDES
# ==========================================
def add_dual_viz_slide(slide, title, slide_num, img_left, cap_left, bullets_left, img_right, cap_right, bullets_right):
    add_header(slide, title, slide_num)
    
    # Left Column Container
    left_x = 0.8
    card_w = 5.6
    
    # Add Left Image
    img_l_path = os.path.join(IMG_DIR, img_left)
    if os.path.exists(img_l_path):
        slide.shapes.add_picture(img_l_path, Inches(left_x), Inches(1.4), width=Inches(card_w), height=Inches(3.4))
    
    # Left Caption & Bullets
    lbox = slide.shapes.add_textbox(Inches(left_x), Inches(4.85), Inches(card_w), Inches(2.2))
    ltf = lbox.text_frame
    ltf.word_wrap = True
    lp1 = ltf.paragraphs[0]
    lp1.text = cap_left
    lp1.font.name = FONT_FAMILY
    lp1.font.size = Pt(11)
    lp1.font.bold = True
    lp1.font.italic = True
    lp1.font.color.rgb = PURPLE_PRIMARY
    lp1.space_after = Pt(4)
    for b in bullets_left:
        bp = ltf.add_paragraph()
        bp.text = "• " + b
        bp.font.name = FONT_FAMILY
        bp.font.size = Pt(11)
        bp.font.color.rgb = TEXT_DARK
        bp.space_after = Pt(2)
        
    # Right Column Container
    right_x = 6.9
    img_r_path = os.path.join(IMG_DIR, img_right)
    if os.path.exists(img_r_path):
        slide.shapes.add_picture(img_r_path, Inches(right_x), Inches(1.4), width=Inches(card_w), height=Inches(3.4))
    
    # Right Caption & Bullets
    rbox = slide.shapes.add_textbox(Inches(right_x), Inches(4.85), Inches(card_w), Inches(2.2))
    rtf = rbox.text_frame
    rtf.word_wrap = True
    rp1 = rtf.paragraphs[0]
    rp1.text = cap_right
    rp1.font.name = FONT_FAMILY
    rp1.font.size = Pt(11)
    rp1.font.bold = True
    rp1.font.italic = True
    rp1.font.color.rgb = PURPLE_PRIMARY
    rp1.space_after = Pt(4)
    for b in bullets_right:
        bp = rtf.add_paragraph()
        bp.text = "• " + b
        bp.font.name = FONT_FAMILY
        bp.font.size = Pt(11)
        bp.font.color.rgb = TEXT_DARK
        bp.space_after = Pt(2)

# ==========================================
# SLIDE 12: DATA VISUALIZATION: AQI DISTRIBUTION & CATEGORIES
# ==========================================
s12 = prs.slides.add_slide(BLANK_LAYOUT)
add_dual_viz_slide(
    s12,
    "Data Visualization: AQI Distribution & Categories",
    12,
    "01_aqi_histogram.png",
    "Figure 10.1 – National Distribution of AQI Values",
    [
        "Right-skewed distribution with overall mean of 111.9 and median of 94.0.",
        "Heavy concentration between 50 and 150; long hazardous tail extending to 500."
    ],
    "02_aqi_category_bar.png",
    "Figure 10.2 – Frequency of AQI Category Classifications",
    [
        "Moderate (101–200) and Satisfactory (51–100) constitute the majority of observations.",
        "Over 22,000+ exposure events fall in the dangerous Very Poor and Severe bands."
    ]
)

# ==========================================
# SLIDE 13: DATA VISUALIZATION: SPATIAL & SEASONAL PATTERNS
# ==========================================
s13 = prs.slides.add_slide(BLANK_LAYOUT)
add_dual_viz_slide(
    s13,
    "Data Visualization: Spatial & Seasonal Patterns",
    13,
    "03_top_states_bar.png",
    "Figure 10.3 – Top Most and Least Polluted States",
    [
        "Delhi (Mean: 206.4) and Jharkhand (164.9) exhibit the highest national pollution.",
        "Mizoram (47.2), Sikkim (53.7) and Tamil Nadu (67.8) maintain the cleanest baselines."
    ],
    "04_monthly_trend.png",
    "Figure 10.4 – Monthly Seasonality & Inversion Spikes",
    [
        "Severe winter peaks in November (161.0) and January (152.0) due to meteorological inversion.",
        "Significant monsoon washout drop in July (62.8) and August (65.5)."
    ]
)

# ==========================================
# SLIDE 14: DATA VISUALIZATION: POLLUTANTS & IMPACT
# ==========================================
s14 = prs.slides.add_slide(BLANK_LAYOUT)
add_dual_viz_slide(
    s14,
    "Data Visualization: Pollutants & Impact",
    14,
    "07_pollutant_distribution.png",
    "Figure 11.1 – Proportion of Dominant Prominent Pollutants",
    [
        "PM10 (47.1%) and PM2.5 (25.3%) drive over 72% of all daily primary index triggers.",
        "Gaseous pollutants (NO₂, SO₂, CO) rarely act as the sole dominant trigger."
    ],
    "08_aqi_by_pollutant.png",
    "Figure 11.2 – Average AQI Impact by Dominant Pollutant",
    [
        "PM2.5-dominant days exhibit the highest severity, averaging an AQI of 168.0.",
        "Combined PM2.5 + PM10 events average 130.0, far exceeding single PM10 (93.7)."
    ]
)

# ==========================================
# SLIDE 15: DATA VISUALIZATION: HEATMAP & FEATURE IMPORTANCE
# ==========================================
s15 = prs.slides.add_slide(BLANK_LAYOUT)
add_dual_viz_slide(
    s15,
    "Data Visualization: Heatmap & Feature Importance",
    15,
    "12_heatmap_state_month.png",
    "Figure 11.3 – Spatio-Temporal State vs Month Heatmap",
    [
        "Intense red clustering across northern states (Delhi, UP, Haryana) from Oct to Feb.",
        "Consistent green (clean) baselines maintained year-round in coastal and north-eastern states."
    ],
    "14_rf_variable_importance.png",
    "Figure 11.4 – Random Forest Variable Importance (%IncMSE)",
    [
        "pollutant_score (%IncMSE: 168.0) and state_mean_aqi (135.1) are the dominant drivers.",
        "Seasonal indicators (is_monsoon, is_winter) provide critical non-linear split power."
    ]
)

# ==========================================
# SLIDE 16: RESULTS: DATA PREPARATION & FINDINGS
# ==========================================
s16 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s16, "Results: Data Preparation & Key Findings", 16)

# Top 4 Stat Cards
res_kpis = [
    ("235,785", "total records analyzed"),
    ("32", "states & UTs mapped"),
    ("55.5%", "CPCB compliant days"),
    ("17.8%", "WHO compliant days")
]
for i, (val, label) in enumerate(res_kpis):
    left = 0.8 + i * 2.95
    create_card(s16, left, 1.4, 2.7, 1.4)
    box = s16.shapes.add_textbox(Inches(left), Inches(1.5), Inches(2.7), Inches(1.2))
    tf = box.text_frame
    p1 = tf.paragraphs[0]
    p1.text = val
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = PURPLE_PRIMARY
    p1.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED
    p2.alignment = PP_ALIGN.CENTER

# 4 Key Findings
f_box = s16.shapes.add_textbox(Inches(0.8), Inches(3.1), Inches(11.733), Inches(3.9))
ftf = f_box.text_frame
ftf.word_wrap = True

findings = [
    ("Extreme Seasonal Disparity", "Winter months (Nov–Jan) experience 2.5x higher AQI than monsoon months (Jul–Aug), driven by boundary layer compression and stubble burning."),
    ("Particulate Matter Toxicity", "PM2.5-driven days average 168.0 AQI versus 93.7 for PM10, proving fine particulate matter is the primary driver of hazardous health spikes."),
    ("Severe Geographical Clustering", "North-Central states exceed national safety limits on over 68% of days, whereas North-Eastern states achieve >90% WHO compliance."),
    ("Monitoring Station Density Effect", "Cities with >10 stations capture 18% higher peak anomalies, proving dense monitoring is essential for timely urban alerts.")
]
for title, desc in findings:
    p1 = ftf.add_paragraph() if ftf.paragraphs[0].text else ftf.paragraphs[0]
    p1.text = title + ":"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = PURPLE_PRIMARY
    
    p2 = ftf.add_paragraph()
    p2.text = desc
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_DARK
    p2.space_after = Pt(8)

# ==========================================
# SLIDE 17: RESULTS: HYPOTHESIS TESTING
# ==========================================
s17 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s17, "Results: Hypothesis Testing", 17)

# Left Card: Welch's t-test
create_card(s17, 0.8, 1.5, 5.7, 4.4)
t_box = s17.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(5.3), Inches(4.1))
ttf = t_box.text_frame
ttf.word_wrap = True
tp1 = ttf.paragraphs[0]
tp1.text = "Welch's Two-Sample t-test (Delhi vs Jharkhand)"
tp1.font.name = FONT_FAMILY
tp1.font.size = Pt(14)
tp1.font.bold = True
tp1.font.color.rgb = PURPLE_PRIMARY
tp1.space_after = Pt(8)

t_bullets = [
    "t = 11.12,  df = 1826.1",
    "p-value < 2.2e-16 (p < 0.001)",
    "95% Confidence Interval: [34.16, 48.79]",
    "Delhi Mean AQI: 206.42 | Jharkhand Mean: 164.94",
    "Null Hypothesis Rejected.",
    "Delhi has a statistically significant higher pollution level (+41.48 AQI) than Jharkhand."
]
for b in t_bullets:
    p = ttf.add_paragraph()
    p.text = "• " + b
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(4)

# Right Card: One-Way ANOVA
create_card(s17, 6.8, 1.5, 5.7, 4.4)
a_box = s17.shapes.add_textbox(Inches(7.0), Inches(1.65), Inches(5.3), Inches(4.1))
atf = a_box.text_frame
atf.word_wrap = True
ap1 = atf.paragraphs[0]
ap1.text = "One-Way ANOVA (3 Factorial Tests)"
ap1.font.name = FONT_FAMILY
ap1.font.size = Pt(14)
ap1.font.bold = True
ap1.font.color.rgb = PURPLE_PRIMARY
ap1.space_after = Pt(8)

a_bullets = [
    "State Effect: F(31, 235753) = 1,674,  p < 2e-16",
    "Monthly Seasonality: F(11, 235773) = 5,403,  p < 2e-16",
    "Pollutant Type: F(48, 235736) = 1,644,  p < 2e-16",
    "All 3 Null Hypotheses Rejected (p < 0.001).",
    "Geography, seasonal weather, and pollutant chemical type all fundamentally alter AQI variance."
]
for b in a_bullets:
    p = atf.add_paragraph()
    p.text = "• " + b
    p.font.name = FONT_FAMILY
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(4)

# Bottom Summary Card
create_card(s17, 0.8, 6.1, 11.7, 0.85, bg_color=RGBColor(255, 255, 255), border_color=PURPLE_PRIMARY)
sbox = s17.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.3), Inches(0.75))
sp = sbox.text_frame.paragraphs[0]
sp.text = "In simple words: All differences are real and statistically proven (p < 0.001). Location, season, and pollutant mix are definitive drivers of air quality."
sp.font.name = FONT_FAMILY
sp.font.size = Pt(12)
sp.font.bold = True
sp.font.color.rgb = PURPLE_PRIMARY

# ==========================================
# SLIDE 18: RESULTS: MODEL PERFORMANCE
# ==========================================
s18 = prs.slides.add_slide(BLANK_LAYOUT)
add_header(s18, "Results: Model Performance", 18)

# Left: Performance Table
m_table_shape = s18.shapes.add_table(3, 4, Inches(0.8), Inches(1.5), Inches(6.0), Inches(1.8))
mt = m_table_shape.table
mt.columns[0].width = Inches(2.2)
mt.columns[1].width = Inches(1.2)
mt.columns[2].width = Inches(1.2)
mt.columns[3].width = Inches(1.4)

for c, h in enumerate(["Model", "MAE", "RMSE", "R²"]):
    cell = mt.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = TABLE_HEADER
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

m_data = [
    ("Linear Regression", "41.4362", "55.6092", "0.3901"),
    ("Random Forest", "34.4654", "48.2871", "0.5412")
]
for r, row in enumerate(m_data):
    for c, val in enumerate(row):
        cell = mt.cell(r + 1, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_ALT if r == 1 else RGBColor(255, 255, 255)
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = FONT_FAMILY
        p.font.size = Pt(11)
        p.font.bold = (r == 1)
        p.font.color.rgb = PURPLE_PRIMARY if r == 1 else TEXT_DARK
        p.alignment = PP_ALIGN.CENTER

# Right: Actual vs Predicted Image
img_rf_path = os.path.join(IMG_DIR, "16_rf_actual_vs_predicted.png")
if os.path.exists(img_rf_path):
    s18.shapes.add_picture(img_rf_path, Inches(7.1), Inches(1.5), width=Inches(5.4), height=Inches(3.3))

# Left Bullets below table
mb_box = s18.shapes.add_textbox(Inches(0.8), Inches(3.6), Inches(6.0), Inches(3.4))
mtf = mb_box.text_frame
mtf.word_wrap = True

m_bullets = [
    "Random Forest achieves superior performance: lowest MAE (34.46) and RMSE (48.28), with highest R² (0.5412).",
    "Random Forest explains 54.1% of the total variance across 235,785 multi-sensor observations.",
    "Key drivers: pollutant_score, state_mean_aqi, and seasonal indicator (is_winter / is_monsoon).",
    "Model embedded as reactive engine in R Shiny dashboard and served via Plumber REST API."
]
for b in m_bullets:
    p = mtf.add_paragraph() if mtf.paragraphs[0].text else mtf.paragraphs[0]
    p.text = "• " + b
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_after = Pt(6)

# ==========================================
# SLIDE 19: THANK YOU
# ==========================================
s19 = prs.slides.add_slide(BLANK_LAYOUT)

# Purple Thank You Text
t_box = s19.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(2.0))
ttf = t_box.text_frame
tp1 = ttf.paragraphs[0]
tp1.text = "THANK YOU"
tp1.font.name = FONT_FAMILY
tp1.font.size = Pt(44)
tp1.font.bold = True
tp1.font.color.rgb = PURPLE_PRIMARY
tp1.alignment = PP_ALIGN.CENTER

tp2 = ttf.add_paragraph()
tp2.text = "Any Questions?"
tp2.font.name = FONT_FAMILY
tp2.font.size = Pt(20)
tp2.font.italic = True
tp2.font.color.rgb = TEXT_MUTED
tp2.alignment = PP_ALIGN.CENTER

# Contact Details Card
create_card(s19, 3.5, 4.6, 6.333, 1.6)
c_box = s19.shapes.add_textbox(Inches(3.6), Inches(4.75), Inches(6.133), Inches(1.3))
ctf = c_box.text_frame
cp1 = ctf.paragraphs[0]
cp1.text = "GitHub: github.com/GokulkrishnaR31/AQI-ANALYSIS-R"
cp1.font.name = FONT_FAMILY
cp1.font.size = Pt(13)
cp1.font.bold = True
cp1.font.color.rgb = PURPLE_PRIMARY
cp1.alignment = PP_ALIGN.CENTER

cp2 = ctf.add_paragraph()
cp2.text = "Email: [your-email@rajalakshmi.edu.in]\nDepartment of Information Technology, REC"
cp2.font.name = FONT_FAMILY
cp2.font.size = Pt(12)
cp2.font.color.rgb = TEXT_DARK
cp2.alignment = PP_ALIGN.CENTER

# Save the presentation
output_pptx = r"g:\rproject\AQI_Data_Science_Presentation.pptx"
prs.save(output_pptx)
print(f"Presentation successfully saved to: {output_pptx}")
